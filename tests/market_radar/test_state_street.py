"""State Street 当前 SPY 持仓来源适配测试。"""

from __future__ import annotations

import asyncio
from http.client import IncompleteRead
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from trading_assistant.market_radar.state_street import (
    STATE_STREET_SPY_HOLDINGS_URL,
    STATE_STREET_SPY_SOURCE,
    StateStreetSpyHoldingsSource,
    download_state_street_holdings,
    parse_state_street_spy_holdings,
)

HEADERS = (
    "Name",
    "Ticker",
    "Identifier",
    "SEDOL",
    "Weight",
    "Sector",
    "Shares Held",
    "Local Currency",
)


def _letters(value: int) -> str:
    result = ""
    current = value
    while True:
        result = chr(ord("A") + current % 26) + result
        current = current // 26 - 1
        if current < 0:
            return result


def _workbook(*, member_count: int = 450) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "holdings"
    sheet.append(["Fund Name:", "State Street SPDR S&P 500 ETF Trust"])
    sheet.append(["Ticker Symbol:", "SPY"])
    sheet.append(["Holdings:", "As of 01-Sep-2026"])
    sheet.append([])
    sheet.append(list(HEADERS))
    symbols = ["AAPL", "BF.B", "BRK.B"]
    symbols.extend(f"Z{_letters(index)}" for index in range(member_count - len(symbols)))
    for index, symbol in enumerate(symbols):
        sheet.append([f"Company {index}", symbol, f"ID{index}", "-", 1, "-", 100, "USD"])
    sheet.append(["US DOLLAR", "-", "USD", "-", 0, "-", 1, "USD"])
    sheet.append(["CONTRA", "2602335D", "ID", "-", 0, "-", 1, "USD"])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _mutate(payload: bytes, *, cell: str, value: object) -> bytes:
    workbook = load_workbook(BytesIO(payload))
    sheet = workbook["holdings"]
    sheet[cell] = value
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def test_parses_strict_snapshot_filters_non_equities_and_maps_class_shares() -> None:
    membership = parse_state_street_spy_holdings(_workbook())

    assert membership.source == STATE_STREET_SPY_SOURCE
    assert membership.membership_date.isoformat() == "2026-09-01"
    assert len(membership.members) == 450
    by_source = {member.source_symbol: member for member in membership.members}
    assert by_source["AAPL"].instrument_id == "AAPL.US"
    assert by_source["BF.B"].data_symbol == "BF-B.US"
    assert by_source["BRK.B"].instrument_id == "BRK-B.US"
    assert "-" not in by_source
    assert "2602335D" not in by_source
    assert tuple(member.instrument_id for member in membership.members) == tuple(
        sorted(member.instrument_id for member in membership.members)
    )


@pytest.mark.parametrize(
    ("cell", "value", "message"),
    [
        ("B3", "2026-09-01", "membership date"),
        ("A5", "Security Name", "columns"),
        ("B2", "QQQ", "identity"),
        ("H6", "CAD", "not USD"),
        ("B6", "ABC.C", "unsupported class-share"),
        ("B6", "aapl", "uppercase"),
    ],
)
def test_workbook_drift_and_unsupported_members_fail_closed(
    cell: str,
    value: object,
    message: str,
) -> None:
    with pytest.raises(ValueError, match=message):
        parse_state_street_spy_holdings(_mutate(_workbook(), cell=cell, value=value))


def test_duplicate_and_implausible_member_count_fail_closed() -> None:
    payload = _mutate(_workbook(), cell="B7", value="AAPL")
    with pytest.raises(ValueError, match="duplicate"):
        parse_state_street_spy_holdings(payload)
    with pytest.raises(ValueError, match="outside"):
        parse_state_street_spy_holdings(_workbook(member_count=449))


def test_invalid_payload_and_missing_sheet_fail_closed() -> None:
    with pytest.raises(ValueError, match="valid XLSX"):
        parse_state_street_spy_holdings(b"not a workbook")
    workbook = Workbook()
    workbook.active.title = "other"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    with pytest.raises(ValueError, match="holdings sheet"):
        parse_state_street_spy_holdings(buffer.getvalue())


def test_source_only_transports_and_returns_neutral_snapshot() -> None:
    captured: dict[str, object] = {}

    def transport(url: str, timeout_seconds: int) -> bytes:
        captured["url"] = url
        captured["timeout"] = timeout_seconds
        return _workbook()

    source = StateStreetSpyHoldingsSource(
        request_timeout_seconds=12,
        transport=transport,
    )
    result = asyncio.run(source.fetch_current_membership())

    assert result.source == STATE_STREET_SPY_SOURCE
    assert captured == {"url": STATE_STREET_SPY_HOLDINGS_URL, "timeout": 12}


def test_source_configuration_is_restricted() -> None:
    with pytest.raises(ValueError, match="positive"):
        StateStreetSpyHoldingsSource(request_timeout_seconds=0)
    with pytest.raises(ValueError, match="HTTPS"):
        StateStreetSpyHoldingsSource(url="http://example.test/holdings.xlsx")


def test_holdings_download_converts_truncated_response(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class TruncatedResponse:
        def __enter__(self) -> TruncatedResponse:
            return self

        def __exit__(self, *_args: object) -> None:
            return None

        def read(self, _size: int) -> bytes:
            raise IncompleteRead(b"partial", 10)

    monkeypatch.setattr(
        "trading_assistant.market_radar.state_street.urlopen",
        lambda *_args, **_kwargs: TruncatedResponse(),
    )
    with pytest.raises(ConnectionError, match="IncompleteRead"):
        download_state_street_holdings(STATE_STREET_SPY_HOLDINGS_URL, 30)
