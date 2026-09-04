"""State Street SPY 每日持仓到当前市场成员契约的适配。"""

from __future__ import annotations

import asyncio
import re
from collections.abc import Callable
from datetime import date, datetime
from http.client import HTTPException, HTTPResponse
from io import BytesIO
from typing import cast
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from trading_assistant.market_radar.membership import (
    CurrentMarketMember,
    CurrentMarketMembership,
)

STATE_STREET_SPY_HOLDINGS_URL = (
    "https://www.ssga.com/library-content/products/fund-data/etfs/us/holdings-daily-us-en-spy.xlsx"
)
STATE_STREET_SPY_SOURCE = "state_street_spy_holdings"

HoldingsTransport = Callable[[str, int], bytes]

_EXPECTED_HEADERS = (
    "Name",
    "Ticker",
    "Identifier",
    "SEDOL",
    "Weight",
    "Sector",
    "Shares Held",
    "Local Currency",
)
_EQUITY_TICKER = re.compile(r"[A-Z]{1,5}(?:\.[A-Z])?")
_CLASS_SHARE_SYMBOLS = {
    "BF.B": "BF-B",
    "BRK.B": "BRK-B",
}
_MIN_EXPECTED_MEMBERS = 450
_MAX_EXPECTED_MEMBERS = 550
_MAX_WORKBOOK_BYTES = 10 * 1024 * 1024


def download_state_street_holdings(url: str, timeout_seconds: int) -> bytes:
    """从固定 HTTPS 地址下载工作簿并隐藏传输细节。"""
    request = Request(  # noqa: S310
        url,
        headers={
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "User-Agent": "HeyBoss-trading-assistant/0.1",
        },
        method="GET",
    )
    try:
        response = cast(HTTPResponse, urlopen(request, timeout=timeout_seconds))  # noqa: S310
        with response:
            return response.read(_MAX_WORKBOOK_BYTES + 1)
    except HTTPError as exc:
        if exc.code == 429 or exc.code >= 500:
            raise RuntimeError(
                f"State Street holdings temporary HTTP failure: status={exc.code}"
            ) from None
        raise ValueError(f"State Street holdings request was rejected: status={exc.code}") from None
    except URLError as exc:
        raise ConnectionError(
            f"State Street holdings connection failed: {type(exc.reason).__name__}"
        ) from None
    except (HTTPException, OSError) as exc:
        raise ConnectionError(
            f"State Street holdings connection interrupted: {type(exc).__name__}"
        ) from None


def _text(value: object, *, field: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"State Street holdings has invalid {field}")
    return value.strip()


def _membership_date(value: object) -> date:
    text = _text(value, field="membership date")
    try:
        return datetime.strptime(text, "As of %d-%b-%Y").date()
    except ValueError as exc:
        raise ValueError("State Street holdings has invalid membership date") from exc


def _eodhd_symbol(source_symbol: str) -> str:
    if "." not in source_symbol:
        return source_symbol
    try:
        return _CLASS_SHARE_SYMBOLS[source_symbol]
    except KeyError as exc:
        raise ValueError(
            f"State Street holdings has unsupported class-share ticker: {source_symbol}"
        ) from exc


def parse_state_street_spy_holdings(payload: bytes) -> CurrentMarketMembership:
    """严格解析官方工作簿, 只输出供应商无关成员模型。"""
    if not payload or len(payload) > _MAX_WORKBOOK_BYTES:
        raise ValueError("State Street holdings workbook size is invalid")
    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("State Street holdings response is not a valid XLSX workbook") from exc
    try:
        if "holdings" not in workbook.sheetnames:
            raise ValueError("State Street holdings workbook is missing the holdings sheet")
        sheet = workbook["holdings"]
        if (
            sheet["A1"].value != "Fund Name:"
            or sheet["A2"].value != "Ticker Symbol:"
            or sheet["B2"].value != "SPY"
            or sheet["A3"].value != "Holdings:"
        ):
            raise ValueError("State Street holdings workbook identity is invalid")
        membership_date = _membership_date(sheet["B3"].value)
        headers = tuple(sheet.cell(row=5, column=index).value for index in range(1, 9))
        if headers != _EXPECTED_HEADERS:
            raise ValueError("State Street holdings columns have changed")

        members: list[CurrentMarketMember] = []
        for row in sheet.iter_rows(min_row=6, max_col=8, values_only=True):
            raw_symbol = row[1]
            if not isinstance(raw_symbol, str):
                continue
            source_symbol = raw_symbol.strip()
            if source_symbol != source_symbol.upper():
                raise ValueError("State Street holdings ticker must be uppercase")
            if _EQUITY_TICKER.fullmatch(source_symbol) is None:
                continue
            _text(row[0], field=f"name for {source_symbol}")
            currency = _text(row[7], field=f"currency for {source_symbol}")
            if currency != "USD":
                raise ValueError(
                    f"State Street holdings member is not USD denominated: {source_symbol}"
                )
            data_root = _eodhd_symbol(source_symbol)
            members.append(
                CurrentMarketMember(
                    source_symbol=source_symbol,
                    instrument_id=f"{data_root}.US",
                    data_symbol=f"{data_root}.US",
                )
            )
    finally:
        workbook.close()

    members.sort(key=lambda item: item.instrument_id)
    if not _MIN_EXPECTED_MEMBERS <= len(members) <= _MAX_EXPECTED_MEMBERS:
        raise ValueError("State Street holdings equity member count is outside the expected range")
    return CurrentMarketMembership(
        source=STATE_STREET_SPY_SOURCE,
        membership_date=membership_date,
        members=tuple(members),
    )


class StateStreetSpyHoldingsSource:
    """首个当前成员来源; 其供应商细节不会越过成员契约。"""

    def __init__(
        self,
        *,
        request_timeout_seconds: int = 30,
        transport: HoldingsTransport = download_state_street_holdings,
        url: str = STATE_STREET_SPY_HOLDINGS_URL,
    ) -> None:
        if request_timeout_seconds < 1:
            raise ValueError("State Street request timeout must be positive")
        if not url.startswith("https://"):
            raise ValueError("State Street holdings URL must use HTTPS")
        self._request_timeout_seconds = request_timeout_seconds
        self._transport = transport
        self._url = url

    async def fetch_current_membership(self) -> CurrentMarketMembership:
        """下载到内存后解析; 不落地供应商原始工作簿。"""
        payload = await asyncio.to_thread(
            self._transport,
            self._url,
            self._request_timeout_seconds,
        )
        return parse_state_street_spy_holdings(payload)
